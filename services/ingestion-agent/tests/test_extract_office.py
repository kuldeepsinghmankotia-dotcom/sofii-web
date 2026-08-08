import io
from unittest.mock import AsyncMock, patch

import pandas as pd
import pytest
from openpyxl import Workbook
from pptx import Presentation

from app.graph.nodes.extract_pptx import extract_pptx_node
from app.graph.nodes.extract_xlsx import extract_xlsx_node
from app.graph.state import IngestionState


def make_state() -> IngestionState:
    return IngestionState(
        job_id="job-1",
        storage_path="uploads/user-1/file",
        filename="file",
        mime_type="application/octet-stream",
        user_id="user-1",
    )


def make_pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Report"
    body = slide.placeholders[1]
    body.text_frame.text = "Revenue grew 12% year over year."

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "amount"])
    sheet.append(["widget", 42])
    sheet.append(["gadget", 7])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_extract_pptx_node_pulls_slide_text() -> None:
    with patch(
        "app.graph.nodes.extract_pptx.download_file_bytes",
        new=AsyncMock(return_value=make_pptx_bytes()),
    ):
        result = await extract_pptx_node(make_state())

    assert result["modality"] == "prose"
    assert "Quarterly Report" in result["extracted_text"]
    assert "Revenue grew 12%" in result["extracted_text"]
    assert "Slide 1" in result["extracted_text"]


@pytest.mark.asyncio
async def test_extract_xlsx_node_produces_parseable_csv() -> None:
    with patch(
        "app.graph.nodes.extract_xlsx.download_file_bytes",
        new=AsyncMock(return_value=make_xlsx_bytes()),
    ):
        result = await extract_xlsx_node(make_state())

    assert result["modality"] == "tabular"
    # Round-trips through pandas the same way chunk_node's tabular path
    # will - proves the CSV text is actually well-formed, not just non-empty.
    df = pd.read_csv(io.StringIO(result["extracted_text"]))
    assert list(df["name"]) == ["widget", "gadget"]
    assert list(df["amount"]) == [42, 7]


@pytest.mark.asyncio
async def test_extract_xlsx_node_rejects_empty_sheet() -> None:
    workbook = Workbook()
    buffer = io.BytesIO()
    workbook.save(buffer)

    with patch(
        "app.graph.nodes.extract_xlsx.download_file_bytes",
        new=AsyncMock(return_value=buffer.getvalue()),
    ):
        with pytest.raises(ValueError, match="no data rows"):
            await extract_xlsx_node(make_state())
