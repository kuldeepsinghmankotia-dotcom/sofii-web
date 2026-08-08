import io

from pptx import Presentation
from pydantic import BaseModel, Field

from app.graph.nodes._storage import download_file_bytes
from app.graph.state import IngestionState


class ExtractedTextResult(BaseModel):
    text: str = Field(min_length=1)


async def extract_pptx_node(state: IngestionState) -> dict:
    raw = await download_file_bytes(state)
    presentation = Presentation(io.BytesIO(raw))

    slides_text: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        # Text-frame shapes only, matching extract_docx_node's scope
        # (paragraph text only) — speaker notes and embedded objects
        # (charts, tables-as-images) aren't walked separately.
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if texts:
            slides_text.append(f"Slide {index}:\n" + "\n".join(texts))

    text = "\n\n".join(slides_text)

    result = ExtractedTextResult(text=text)
    return {"extracted_text": result.text, "modality": "prose"}
